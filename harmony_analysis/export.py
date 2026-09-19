"""
export.py
---------
Writes per_trial_metrics.csv, aggregate_results.csv, paper_metrics.csv,
validation_report.csv, event_diagnostics.csv and the multi-sheet
report_summary.xlsx (spec section 6).
"""
from __future__ import annotations

from pathlib import Path
from typing import Dict, Optional

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet


def write_csvs(output_dir: Path,
               per_trial: pd.DataFrame,
               aggregate: pd.DataFrame,
               paper_metrics: pd.DataFrame,
               validation_report: pd.DataFrame,
               event_diagnostics: pd.DataFrame) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    per_trial.to_csv(output_dir / "per_trial_metrics.csv", index=False)
    aggregate.to_csv(output_dir / "aggregate_results.csv", index=False)
    paper_metrics.to_csv(output_dir / "paper_metrics.csv", index=False)
    validation_report.to_csv(output_dir / "validation_report.csv", index=False)
    event_diagnostics.to_csv(output_dir / "event_diagnostics.csv", index=False)


def _write_sheet(wb: Workbook, name: str, df: pd.DataFrame,
                  highlight_invalid: bool = False) -> None:
    ws: Worksheet = wb.create_sheet(title=name[:31])
    if df is None or df.empty:
        ws.append(["(no data)"])
        return

    ws.append(list(df.columns))
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for _, row in df.iterrows():
        ws.append([_excel_safe(v) for v in row.tolist()])

    # freeze header + autofilter
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # sensible column widths
    for col_idx, col_name in enumerate(df.columns, start=1):
        try:
            max_len = max(
                [len(str(col_name))] +
                [len(str(v)) for v in df[col_name].astype(str).tolist()[:500]]
            )
        except Exception:
            max_len = 12
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max(max_len + 2, 8), 45)

    if highlight_invalid and "validation_status" in df.columns:
        fill_invalid = PatternFill(start_color="F8CBAD", end_color="F8CBAD", fill_type="solid")
        fill_warn = PatternFill(start_color="FFF2CC", end_color="FFF2CC", fill_type="solid")
        status_col_idx = list(df.columns).index("validation_status") + 1
        for r in range(2, ws.max_row + 1):
            status_val = ws.cell(row=r, column=status_col_idx).value
            if status_val == "INVALID":
                for c in range(1, ws.max_column + 1):
                    ws.cell(row=r, column=c).fill = fill_invalid
            elif status_val == "VALID_WITH_WARNINGS":
                for c in range(1, ws.max_column + 1):
                    ws.cell(row=r, column=c).fill = fill_warn


def _excel_safe(v):
    if isinstance(v, float) and (v != v):  # NaN
        return None
    return v


def write_excel_report(output_path: Path,
                        per_trial: pd.DataFrame,
                        aggregate: pd.DataFrame,
                        paper_metrics: pd.DataFrame,
                        validation_report: pd.DataFrame,
                        readme_lines: list[str],
                        source_transitions: Optional[pd.DataFrame] = None,
                        gps_stats: Optional[pd.DataFrame] = None,
                        pdr_stats: Optional[pd.DataFrame] = None,
                        vps_stats: Optional[pd.DataFrame] = None) -> None:
    wb = Workbook()
    wb.remove(wb.active)

    readme_df = pd.DataFrame({"README": readme_lines})
    _write_sheet(wb, "README", readme_df)
    _write_sheet(wb, "Per Trial", per_trial, highlight_invalid=True)
    _write_sheet(wb, "Aggregate", aggregate)
    _write_sheet(wb, "Paper Metrics", paper_metrics)
    _write_sheet(wb, "Validation", validation_report, highlight_invalid=True)

    switch_cols = [c for c in per_trial.columns if c.startswith("switch_count_")] + \
                  ["session_id", "harmony_version", "direction", "source_transitions_total",
                   "unnecessary_source_switches", "oscillation_episode_count"]
    switch_cols = [c for c in switch_cols if c in per_trial.columns]
    _write_sheet(wb, "Source Transitions", per_trial[switch_cols] if switch_cols else pd.DataFrame())

    gps_cols = [c for c in per_trial.columns if c.startswith("gps_")] + ["session_id", "harmony_version", "direction"]
    gps_cols = [c for c in dict.fromkeys(gps_cols) if c in per_trial.columns]
    _write_sheet(wb, "GPS Statistics", per_trial[gps_cols] if gps_cols else pd.DataFrame())

    pdr_cols = [c for c in per_trial.columns if c.startswith("pdr_")] + ["session_id", "harmony_version", "direction"]
    pdr_cols = [c for c in dict.fromkeys(pdr_cols) if c in per_trial.columns]
    _write_sheet(wb, "PDR Statistics", per_trial[pdr_cols] if pdr_cols else pd.DataFrame())

    vps_cols = [c for c in per_trial.columns if c.startswith("vps_")] + ["session_id", "harmony_version", "direction"]
    vps_cols = [c for c in dict.fromkeys(vps_cols) if c in per_trial.columns]
    _write_sheet(wb, "VPS Statistics", per_trial[vps_cols] if vps_cols else pd.DataFrame())

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
